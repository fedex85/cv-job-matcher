"""Extracción de texto de DOCX con python-docx, con defensas contra zip bombs."""

from __future__ import annotations

import io
import zipfile

from .errors import IngestionError

MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000


def is_docx(data: bytes) -> bool:
    """Comprueba por contenido (no por extensión) que es un .docx."""
    if data[:4] != b"PK\x03\x04":
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return "word/document.xml" in zf.namelist()
    except zipfile.BadZipFile:
        return False


def _check_zip_safety(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise IngestionError("unsafe_docx", "El DOCX contiene demasiados elementos internos.")
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
                raise IngestionError("unsafe_docx", "El DOCX descomprimido es demasiado grande.", 413)
    except zipfile.BadZipFile as exc:
        raise IngestionError("corrupt_docx", "No se pudo leer el DOCX: el archivo parece dañado.") from exc


def parse_docx(data: bytes, *, max_chars: int) -> str:
    _check_zip_safety(data)
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError("corrupt_docx", "No se pudo leer el DOCX: el archivo parece dañado.") from exc

    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))  # quita celdas combinadas repetidas
    text = "\n".join(parts)
    if len(text) > max_chars:
        raise IngestionError(
            "text_too_long", f"El documento supera los {max_chars} caracteres de texto permitidos.", 413
        )
    return text
